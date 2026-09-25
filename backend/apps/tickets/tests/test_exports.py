from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.tickets.models import Ticket
from apps.tickets.permissions import ROLE_DEV, ROLE_PM

from .helpers import make_project, make_routed_category, make_user


class ExportTests(TestCase):
    def setUp(self):
        self.pm = make_user("pm", ROLE_PM)
        self.dev = make_user("dev", ROLE_DEV)
        category, _ = make_routed_category()
        project = make_project()
        for title, priority in (("Нэвтрэх алдаа", Ticket.Priority.HIGH), ("Other", Ticket.Priority.LOW)):
            Ticket.objects.create(
                title=title, description="d", ticket_type=Ticket.TicketType.BUG,
                category=category, project=project, reported_by=self.dev, priority=priority,
            )

    def test_ticket_list_csv_respects_filters_and_has_bom(self):
        self.client.force_login(self.pm)
        response = self.client.get(
            reverse("tickets:ticket_list"), {"priority": Ticket.Priority.HIGH, "export": "csv"}
        )
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("attachment;", response["Content-Disposition"])
        body = response.content.decode("utf-8")
        self.assertTrue(body.startswith("﻿"))
        self.assertIn("Нэвтрэх алдаа", body)
        self.assertNotIn("Other", body)

    def test_ticket_list_xlsx(self):
        self.client.force_login(self.pm)
        response = self.client.get(reverse("tickets:ticket_list"), {"export": "xlsx"})
        ws = load_workbook(BytesIO(response.content)).active
        self.assertEqual(ws.max_row, 3)  # толгой + 2 ticket

    def test_report_xlsx_has_all_sheets(self):
        self.client.force_login(self.pm)
        response = self.client.get(reverse("tickets:reports"), {"days": 7, "export": "xlsx"})
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(
            wb.sheetnames,
            ["Хураангуй", "SLA", "Ажилтан", "Чухлын зэрэг", "Ангилал", "Урсгал", "Ticket-үүд"],
        )
        self.assertIn("report_7d_", response["Content-Disposition"])

    def test_report_csv(self):
        self.client.force_login(self.pm)
        body = self.client.get(reverse("tickets:reports"), {"export": "csv"}).content.decode()
        self.assertIn("# SLA биелэлт", body)

    def test_developer_cannot_export_report(self):
        self.client.force_login(self.dev)
        response = self.client.get(reverse("tickets:reports"), {"export": "xlsx"})
        self.assertEqual(response.status_code, 302)


class TeamLeadScopeTests(TestCase):
    def setUp(self):
        self.admin = make_user("admin", "Admin")
        self.lead = make_user("lead", ROLE_DEV)
        self.dev = make_user("dev", ROLE_DEV)
        project = make_project()
        cat_a, self.team_a = make_routed_category(team_lead=self.lead, name="A")
        cat_b, self.team_b = make_routed_category(name="B")
        for title, cat in (("Ticket-A", cat_a), ("Ticket-B", cat_b)):
            Ticket.objects.create(
                title=title, description="d", ticket_type=Ticket.TicketType.BUG,
                category=cat, project=project, reported_by=self.dev,
            )

    def test_team_lead_redirected_from_global_dashboard(self):
        self.client.force_login(self.lead)
        response = self.client.get(reverse("tickets:reports"))
        self.assertRedirects(response, reverse("tickets:my_team"))

    def test_team_lead_sees_own_team_dashboard_only(self):
        self.client.force_login(self.lead)
        own = self.client.get(reverse("tickets:team_report", args=[self.team_a.pk]))
        self.assertEqual(own.status_code, 200)
        self.assertEqual(own.context["created_count"], 1)
        other = self.client.get(reverse("tickets:team_report", args=[self.team_b.pk]))
        self.assertEqual(other.status_code, 302)

    def test_team_report_export_contains_only_team_tickets(self):
        self.client.force_login(self.lead)
        response = self.client.get(
            reverse("tickets:team_report", args=[self.team_a.pk]), {"export": "xlsx"}
        )
        ws = load_workbook(BytesIO(response.content))["Ticket-үүд"]
        titles = [row[1] for row in ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(titles, ["Ticket-A"])

    def test_team_lead_list_export_limited_to_team(self):
        self.client.force_login(self.lead)
        body = self.client.get(reverse("tickets:ticket_list"), {"export": "csv"}).content.decode()
        self.assertIn("Ticket-A", body)
        self.assertNotIn("Ticket-B", body)

    def test_developer_cannot_export_list(self):
        self.client.force_login(self.dev)
        response = self.client.get(reverse("tickets:ticket_list"), {"export": "csv"})
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("attachment", response.get("Content-Disposition", ""))

    def test_admin_sees_global_and_any_team(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("tickets:reports")).context["created_count"], 2)
        self.assertEqual(
            self.client.get(reverse("tickets:team_report", args=[self.team_b.pk])).status_code, 200
        )

    def test_pm_who_leads_team_still_sees_global_and_other_teams(self):
        pm_lead = make_user("pmlead", ROLE_PM)
        self.team_a.team_lead = pm_lead
        self.team_a.save()
        self.client.force_login(pm_lead)
        self.assertEqual(self.client.get(reverse("tickets:reports")).context["created_count"], 2)
        self.assertEqual(
            self.client.get(reverse("tickets:team_report", args=[self.team_b.pk])).status_code, 200
        )

    def test_admin_dashboard_has_team_switcher(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("tickets:reports"))
        self.assertContains(response, reverse("tickets:team_report", args=[self.team_b.pk]))

    def test_team_lead_has_no_team_switcher(self):
        self.client.force_login(self.lead)
        response = self.client.get(reverse("tickets:team_report", args=[self.team_a.pk]))
        self.assertNotIn("team_choices", response.context)
