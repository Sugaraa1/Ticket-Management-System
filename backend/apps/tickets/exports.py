"""
Тайлан / ticket жагсаалтыг файл болгон татах (Jira, Zendesk-ийн "Export"-тэй адил).

  * CSV  — UTF-8 BOM-той тул Excel-д монгол үсэг зөв харагдана.
  * XLSX — openpyxl; тайлан нь sheet бүрт хэсэг тус бүрээ (Хураангуй, SLA,
           Ажилтан, Урсгал, Задаргаа, Ticket-үүд) агуулна.
PDF-ийг браузерын хэвлэх (print) загвараар гаргадаг тул энд байхгүй.
"""
import csv
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy as _lazy

from .models import Ticket

CSV_CONTENT_TYPE = "text/csv; charset=utf-8"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_FORMATS = ("csv", "xlsx")

TICKET_HEADERS = [
    "ID",
    _lazy("Гарчиг"),
    _lazy("Төрөл"),
    _lazy("Төлөв"),
    _lazy("Чухлын зэрэг"),
    _lazy("Ангилал"),
    _lazy("Төсөл"),
    _lazy("Модуль"),
    _lazy("Баг"),
    _lazy("Хариуцагч"),
    _lazy("Мэдээлсэн"),
    _lazy("Үүсгэсэн"),
    _lazy("Шинэчилсэн"),
    _lazy("Эхний хариу (хугацаа)"),
    _lazy("Эхний хариу өгсөн"),
    _lazy("Шийдвэрлэх хугацаа (SLA)"),
    _lazy("Хугацаа хэтэрсэн"),
]


def _dt(value):
    """Excel/CSV-д timezone-гүй орон нутгийн цаг болгоно."""
    if not value:
        return None
    return timezone.localtime(value).replace(tzinfo=None, microsecond=0)


def _name(user):
    if not user:
        return ""
    return user.get_full_name() or user.username


def ticket_rows(tickets):
    for t in tickets:
        yield [
            t.code,
            t.title,
            t.get_ticket_type_display(),
            t.get_status_display(),
            t.get_priority_display(),
            t.category.name if t.category_id else "",
            t.project.name if t.project_id else "",
            t.module.name if t.module_id else "",
            t.team.name if t.team_id else "",
            _name(t.assigned_to),
            _name(t.reported_by),
            _dt(t.created_at),
            _dt(t.updated_at),
            _dt(t.first_response_due_at),
            _dt(t.first_responded_at),
            _dt(t.sla_due_at),
            _("Тийм") if t.is_overdue else _("Үгүй"),
        ]


def export_filename(base, ext):
    return f"{base}_{timezone.localdate():%Y%m%d}.{ext}"


def _attachment(response, filename):
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _csv_writer(response):
    response.write("\ufeff")  # BOM — Excel UTF-8 гэж танина
    return csv.writer(response)


def _safe_text(value):
    """
    "=", "+", "-", "@"-аар эхэлсэн текстийг Excel томьёо болгож ажиллуулахаас
    сэргийлнэ (CSV/formula injection) — ticket-ийн гарчгийг хэн ч бичиж чадна.
    """
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _csv_value(value):
    return value.strftime("%Y-%m-%d %H:%M") if hasattr(value, "strftime") else _safe_text(value)


# ---------------------------------------------------------------- Ticket list


def tickets_csv(tickets, filename):
    response = HttpResponse(content_type=CSV_CONTENT_TYPE)
    writer = _csv_writer(response)
    writer.writerow(TICKET_HEADERS)
    for row in ticket_rows(tickets):
        writer.writerow([_csv_value(v) for v in row])
    return _attachment(response, filename)


def tickets_xlsx(tickets, filename):
    wb = _workbook()
    ws = wb.active
    ws.title = _("Ticket-үүд")
    _write_table(ws, TICKET_HEADERS, ticket_rows(tickets))
    return _xlsx_response(wb, filename)


# ---------------------------------------------------------------- Report


def _report_meta(report, user):
    if report.get("team"):
        title = _("%(name)s багийн тайлан") % {"name": report["team"].name}
    else:
        title = _("Ticket-ийн тайлан")
    return [
        (_("Тайлан"), title),
        (_("Хугацаа"), _("Сүүлийн %(days)s хоног") % {"days": report["days"]}),
        (_("Гаргасан"), _dt(timezone.now())),
        (_("Гаргасан хэрэглэгч"), _name(user)),
    ]


def _summary_rows(report):
    return [
        (_("Шинээр үүссэн"), report["created_count"]),
        (_("Хаагдсан"), report["closed_count"]),
        (_("Одоо нээлттэй"), report["open_count"]),
        (_("Дундаж шийдвэрлэх хугацаа (цаг)"), report["avg_resolution_hours"]),
    ]


SLA_HEADERS = [
    _lazy("Хэмжүүр"),
    _lazy("Биелсэн"),
    _lazy("Зөрчигдсөн"),
    _lazy("Хүлээгдэж буй"),
    _lazy("Биелэлт %"),
]


def _sla_rows(report):
    for label, key in ((_("Эхний хариу"), "first_response"), (_("Шийдвэрлэлт"), "resolution")):
        s = report[key]
        yield [label, s["met"], s["breached"], s["pending"], s["percent"]]


AGENT_HEADERS = [
    _lazy("Хэрэглэгч"),
    _lazy("Оноогдсон"),
    _lazy("Идэвхтэй"),
    _lazy("Хаасан"),
    _lazy("SLA зөрчсөн"),
    _lazy("Дундаж шийдвэрлэх хугацаа (цаг)"),
]


def _agent_rows(report):
    for row in report["agents"]:
        yield [
            _name(row["user"]), row["assigned"], row["active"], row["closed"],
            row["breached"], row["avg_resolution_hours"],
        ]


TREND_HEADERS = [
    _lazy("Огноо"),
    _lazy("Үүссэн"),
    _lazy("Хаагдсан"),
]
BREAKDOWN_HEADERS = [
    _lazy("Нэр"),
    _lazy("Тоо"),
    _lazy("Хувь %"),
]


def _trend_rows(report):
    for row in report["trend"]:
        yield [row["date"], row["created"], row["closed"]]


def _breakdown_rows(rows):
    for row in rows:
        yield [str(row["label"]), row["count"], row["percent"]]


def report_csv(report, user, filename):
    response = HttpResponse(content_type=CSV_CONTENT_TYPE)
    writer = _csv_writer(response)

    def section(title, headers, rows):
        writer.writerow([])
        writer.writerow([f"# {title}"])
        writer.writerow(headers)
        for row in rows:
            writer.writerow([_csv_value(v) for v in row])

    for key, value in _report_meta(report, user):
        writer.writerow([key, _csv_value(value)])
    section(_("Хураангуй"), [_("Үзүүлэлт"), _("Утга")], _summary_rows(report))
    section(_("SLA биелэлт"), SLA_HEADERS, _sla_rows(report))
    section(_("Ажилтны гүйцэтгэл"), AGENT_HEADERS, _agent_rows(report))
    section(_("Чухлын зэргээр"), BREAKDOWN_HEADERS, _breakdown_rows(report["by_priority"]))
    section(_("Ангиллаар"), BREAKDOWN_HEADERS, _breakdown_rows(report["by_category"]))
    section(_("Өдөр тутмын урсгал"), TREND_HEADERS, _trend_rows(report))
    return _attachment(response, filename)


def report_xlsx(report, tickets, user, filename):
    from openpyxl.styles import Font

    wb = _workbook()
    ws = wb.active
    ws.title = _("Хураангуй")
    ws.append([_report_meta(report, user)[0][1]])
    ws["A1"].font = Font(bold=True, size=14)
    for key, value in _report_meta(report, user)[1:]:
        ws.append([key, value])
    ws.append([])
    start = ws.max_row + 1
    _write_table(ws, [_("Үзүүлэлт"), _("Утга")], _summary_rows(report), start_row=start)

    sheets = [
        ("SLA", SLA_HEADERS, _sla_rows(report)),
        (_("Ажилтан"), AGENT_HEADERS, _agent_rows(report)),
        (_("Чухлын зэрэг"), BREAKDOWN_HEADERS, _breakdown_rows(report["by_priority"])),
        (_("Ангилал"), BREAKDOWN_HEADERS, _breakdown_rows(report["by_category"])),
        (_("Урсгал"), TREND_HEADERS, _trend_rows(report)),
        (_("Ticket-үүд"), TICKET_HEADERS, ticket_rows(tickets)),
    ]
    for title, headers, rows in sheets:
        _write_table(wb.create_sheet(str(title)), headers, rows)
    return _xlsx_response(wb, filename)


# ---------------------------------------------------------------- XLSX helpers


def _workbook():
    from openpyxl import Workbook

    return Workbook()


def _write_table(ws, headers, rows, start_row=1):
    """Толгой мөрийг тодруулж, шүүлтүүр, freeze pane, баганын өргөн тохируулна."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    widths = [len(str(h)) for h in headers]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=str(header))
        cell.font = header_font
        cell.fill = header_fill

    row_idx = start_row
    for row_idx, row in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col, value=_safe_text(value))
            if hasattr(value, "strftime"):
                cell.number_format = "yyyy-mm-dd hh:mm"
                length = 16
            else:
                length = len(str(value)) if value is not None else 0
            widths[col - 1] = max(widths[col - 1], length)

    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = min(width + 2, 60)
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A{start_row}:{last_col}{max(row_idx, start_row)}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _xlsx_response(wb, filename):
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    return _attachment(response, filename)


def report_period_tickets(since, team=None):
    scope = Ticket.objects.filter(team=team) if team else Ticket.objects.all()
    return (
        scope.filter(created_at__gte=since)
        .select_related("category", "project", "module", "team", "assigned_to", "reported_by")
        .order_by("-created_at")
    )
