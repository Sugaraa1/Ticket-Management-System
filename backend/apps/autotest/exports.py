"""Жишээ загвар файл болон ажиллуулалтын үр дүнг Excel эсвэл CSV-ээр гаргах."""
from io import BytesIO

from django.http import FileResponse
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.tickets.exports import XLSX_CONTENT_TYPE
from apps.tickets.exports import _safe_text as _safe  # томьёо (=, +, -, @, tab) escape

from .datafiles import write_csv

CSV_CONTENT_TYPE = "text/csv; charset=utf-8"

HEADER_FONT = Font(bold=True)
FILLS = {
    "pass": PatternFill("solid", fgColor="D1F2DC"),
    "fail": PatternFill("solid", fgColor="F8D4D4"),
    "error": PatternFill("solid", fgColor="FCE9C7"),
    "recorded": PatternFill("solid", fgColor="EEF0F3"),
}


def requested_format(request):
    """?format=csv бол 'csv', бусад үед 'xlsx'."""
    return "csv" if request.GET.get("format") == "csv" else "xlsx"


def file_response(content, filename, fmt):
    """filename өргөтгөлгүй — fmt-ээр нэмнэ. Кирилл нэрийг FileResponse зөв кодчилно."""
    content_type = CSV_CONTENT_TYPE if fmt == "csv" else XLSX_CONTENT_TYPE
    return FileResponse(BytesIO(content), as_attachment=True, filename=f"{filename}.{fmt}", content_type=content_type)


def _response(workbook, filename):
    buffer = BytesIO()
    workbook.save(buffer)
    return file_response(buffer.getvalue(), filename, "xlsx")


def _autosize(sheet):
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 60)


def template_response(fmt="xlsx"):
    rows = [
        [_("Тайлбар"), "email", "password", "phone", _("хүлээгдэх")],
        [_("Зөв бүртгэл"), "test{{random}}@mail.mn", "Pass123!", "99112233", _("амжилттай")],
        [_("Буруу имэйл"), "bat@@mail", "Pass123!", "99112233", _("алдаа")],
        [_("Богино нууц үг"), "bold{{random}}@mail.mn", "12", "99112233", _("алдаа: нууц үг")],
        [_("Утас хоосон"), "dorj{{random}}@mail.mn", "Pass123!", "", _("алдаа")],
    ]
    if fmt == "csv":  # CSV нэг хуудастай — зааврын sheet-гүй
        return file_response(write_csv(rows[0], rows[1:]), "autotest_template", "csv")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _("Өгөгдөл")
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = HEADER_FONT
    _autosize(sheet)

    guide = workbook.create_sheet(_("Заавар"))
    for line in [
        [_("Файлыг хэрхэн бөглөх вэ")],
        [""],
        [_("1. Эхний мөрөнд баганын нэрс байна. Нэрийг хүссэнээрээ өгч болно — сценари тохируулахдаа талбартай холбоно.")],
        [_("2. Мөр бүр нэг тест. Хоосон мөрийг алгасна.")],
        [_("3. 'хүлээгдэх' баганад: амжилттай / алдаа, эсвэл 'алдаа: мессежийн хэсэг' гэж бичнэ.")],
        [_("   Мессеж бичвэл хуудсан дээр тэр текст гарсан эсэхийг бас шалгана.")],
        [_("4. {{random}} — мөр бүрт шинэ санамсаргүй тэмдэгт (давхардахгүй имэйл үүсгэхэд).")],
        [_("   {{row}} — мөрийн дугаар, {{timestamp}} — одоогийн цаг, {{digits}} — мөр бүрт шинэ 6 орон (утас).")],
        [_("   {{run}} — нэг ажиллуулалтын бүх мөрт ижил (давхар бүртгэл шалгахад).")],
        [_("5. 'Тайлбар' багана байвал үр дүнд тестийн нэр болж харагдана.")],
        [_("6. Нэг файлыг олон сценарид (бүртгэл, нэвтрэх ...) дахин ашиглаж болно.")],
    ]:
        guide.append(line)
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 110
    return _response(workbook, "autotest_template")


def run_response(run, results, fmt="xlsx"):
    columns = []
    for result in results:
        for column in result.input_data:
            if column not in columns:
                columns.append(column)
    header = [_("Мөр"), *[_safe(c) for c in columns], _("Бодит үр дүн"), _("Бодит мессеж"), _("Дүн"), _("Хугацаа (ms)"), _("Эцсийн хаяг")]
    rows = [[
        result.row_number,
        *[_safe(result.input_data.get(c, "")) for c in columns],
        result.get_actual_outcome_display() if result.actual_outcome else "",
        _safe(result.actual_message),
        result.get_verdict_display(),
        result.duration_ms,
        _safe(result.final_url),
    ] for result in results]
    filename = f"autotest_run_{run.pk}"
    if fmt == "csv":
        return file_response(write_csv(header, rows), filename, "csv")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _("Үр дүн")
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = HEADER_FONT
    for result, row in zip(results, rows):
        sheet.append(row)
        fill = FILLS.get(result.verdict)
        if fill:
            sheet.cell(row=sheet.max_row, column=len(header) - 2).fill = fill
    _autosize(sheet)
    return _response(workbook, filename)
