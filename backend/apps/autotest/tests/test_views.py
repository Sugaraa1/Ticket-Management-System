import json
from io import BytesIO

from django.core.files.base import ContentFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.autotest.datafiles import read_rows
from apps.autotest.models import DataFile, Page, PageScan, RunResult, Scenario, TestAccount, TestApp, TestRun
from apps.categories.models import Category, Subcategory
from apps.tickets.models import Ticket
from apps.tickets.permissions import ROLE_ADMIN, ROLE_DEV, ROLE_PM, ROLE_QA
from apps.tickets.tests.helpers import make_project, make_routed_category, make_user

from .helpers import REGISTER_FIELDS, TempMediaMixin, csv_upload, make_category, make_setup, xlsx_upload


class PermissionTests(TempMediaMixin, TestCase):
    def setUp(self):
        self.app, self.env, self.scenario, self.data_file = make_setup()

    def test_pm_and_developer_have_no_access(self):
        for username, role in (("pm", ROLE_PM), ("dev", ROLE_DEV)):
            with self.subTest(role=role):
                self.client.force_login(make_user(username, role))
                self.assertNotEqual(self.client.get(reverse("autotest:home")).status_code, 200)
                self.assertNotEqual(
                    self.client.get(reverse("autotest:scenario_detail", args=[self.scenario.pk])).status_code, 200
                )
                self.client.post(reverse("autotest:run_create", args=[self.scenario.pk]),
                                 {"data_file": self.data_file.pk, "environment": self.env.pk})
                self.client.post(reverse("autotest:app_create"), {"category": self.app.category_id, "name": role})
                self.assertFalse(TestRun.objects.exists())
                self.assertFalse(TestApp.objects.filter(name=role).exists())
                self.assertNotContains(self.client.get(reverse("tickets:ticket_list")), reverse("autotest:home"))

    def test_qa_can_register_apps_and_create_scenarios(self):
        self.client.force_login(make_user("qa", ROLE_QA))
        self.client.post(reverse("autotest:app_create"), {"category": self.app.category_id, "name": "Other"})
        self.assertTrue(TestApp.objects.filter(name="Other", category=self.app.category).exists())
        response = self.client.get(reverse("autotest:scenario_create", args=[self.app.pk]))
        self.assertEqual(response.status_code, 200)

    def test_nav_link_visible_to_qa(self):
        self.client.force_login(make_user("qa", ROLE_QA))
        self.assertContains(self.client.get(reverse("tickets:ticket_list")), reverse("autotest:home"))


class DataFileViewTests(TempMediaMixin, TestCase):
    def setUp(self):
        self.category = make_category()
        self.client.force_login(make_user("qa", ROLE_QA))

    def test_upload_reads_columns_and_rows(self):
        upload = xlsx_upload([["email", "password"], ["a@mail.mn", "x"], ["b@mail.mn", "y"]])
        response = self.client.post(reverse("autotest:datafile_create"), {"category": self.category.pk, "file": upload})
        data_file = DataFile.objects.get()
        self.assertRedirects(response, reverse("autotest:datafile_detail", args=[data_file.pk]))
        self.assertEqual(data_file.name, "users")
        self.assertEqual(data_file.columns, ["email", "password"])
        self.assertEqual(data_file.row_count, 2)
        detail = self.client.get(reverse("autotest:datafile_detail", args=[data_file.pk]))
        self.assertContains(detail, "b@mail.mn")

    def test_broken_file_shows_error(self):
        upload = xlsx_upload([["email", "email"], ["a", "b"]])
        response = self.client.post(reverse("autotest:datafile_create"), {"category": self.category.pk, "file": upload})
        self.assertContains(response, "Давхардсан")
        self.assertFalse(DataFile.objects.exists())

    def test_detail_lists_compatible_scenarios(self):
        _app, _env, scenario, data_file = make_setup(category=self.category)
        page = Page.objects.create(app=scenario.app, name="Утас", path="/p")
        Scenario.objects.create(app=scenario.app, name="Утас", page=page,
                                fields=[{"label": "Утас", "selector": "#p", "source": "column", "value": "phone"}])
        response = self.client.get(reverse("autotest:datafile_detail", args=[data_file.pk]))
        self.assertContains(response, "Бэлэн")
        self.assertContains(response, "Дутуу багана:")

    def test_template_download(self):
        response = self.client.get(reverse("autotest:template_download"))
        sheet = load_workbook(BytesIO(b"".join(response.streaming_content))).worksheets[0]
        self.assertEqual(sheet["B1"].value, "email")

    def test_template_download_as_csv(self):
        response = self.client.get(reverse("autotest:template_download"), {"format": "csv"})
        self.assertIn("text/csv", response["Content-Type"])
        columns, rows = read_rows(BytesIO(b"".join(response.streaming_content)), "t.csv")
        self.assertEqual(columns[1], "email")
        self.assertEqual(rows[0][1]["email"], "test{{random}}@mail.mn")

    def test_download_converts_between_formats(self):
        upload = csv_upload("email,password\nа@mail.mn,-5\n")
        self.client.post(reverse("autotest:datafile_create"), {"category": self.category.pk, "file": upload})
        data_file = DataFile.objects.get()
        url = reverse("autotest:datafile_download", args=[data_file.pk])
        for fmt in ("csv", "xlsx"):
            with self.subTest(fmt=fmt):
                response = self.client.get(url, {"format": fmt})
                self.assertIn(f'.{fmt}"', response["Content-Disposition"])
                columns, rows = read_rows(BytesIO(b"".join(response.streaming_content)), f"f.{fmt}")
                self.assertEqual(columns, ["email", "password"])
                self.assertEqual(rows[0][1], {"email": "а@mail.mn", "password": "-5"})

    def test_category_with_tests_cannot_be_deleted(self):
        make_setup(category=self.category)
        self.client.force_login(make_user("admin", ROLE_ADMIN))
        self.client.post(reverse("categories:category_delete", args=[self.category.pk]))
        self.assertTrue(Category.objects.filter(pk=self.category.pk).exists())


class ScenarioViewTests(TempMediaMixin, TestCase):
    def setUp(self):
        self.app, self.env, scenario, _file = make_setup()
        self.page = scenario.page
        self.client.force_login(make_user("qa", ROLE_QA))

    def _post(self, fields, **extra):
        data = {"name": "Нэвтрэх", "page": self.page.pk, "success_mode": "auto", "fields_json": json.dumps(fields)}
        data.update(extra)
        return self.client.post(reverse("autotest:scenario_create", args=[self.app.pk]), data)

    def test_create_saves_field_mapping(self):
        response = self._post(REGISTER_FIELDS, submit_selector="#go", submit_label="Нэвтрэх")
        scenario = Scenario.objects.get(name="Нэвтрэх")
        self.assertRedirects(response, reverse("autotest:scenario_detail", args=[scenario.pk]))
        self.assertEqual([f["source"] for f in scenario.fields], ["column", "column", "check"])
        self.assertEqual(scenario.required_columns(), ["email", "password"])
        self.assertEqual(scenario.secret_columns(), {"password"})

    def test_requires_at_least_one_field(self):
        response = self._post([{"label": "x", "selector": "#x", "source": "skip"}])
        self.assertContains(response, "Дор хаяж нэг талбар")

    def test_column_source_needs_column_name(self):
        response = self._post([{"label": "Имэйл", "selector": "#e", "source": "column", "value": ""}])
        self.assertContains(response, "багана сонгоогүй")

    def test_absolute_url_page_is_rejected(self):
        """Бүтэн URL бичвэл орчны хязгаарлалтыг тойрох тул хуудас болгож хүлээж авахгүй."""
        for path in ("https://evil.example.com/x", "//evil.example.com/x"):
            response = self.client.post(reverse("autotest:page_create", args=[self.app.pk]),
                                        {"page-name": path, "page-path": path})
            self.assertContains(response, "Бүтэн хаяг биш")
        self.assertEqual(self.app.pages.count(), 1)

    def test_page_is_added_and_path_normalized(self):
        self.client.post(reverse("autotest:page_create", args=[self.app.pk]),
                         {"page-name": "Нэвтрэх", "page-path": "login"})
        self.assertEqual(self.app.pages.get(name="Нэвтрэх").path, "/login")
        response = self.client.post(reverse("autotest:page_create", args=[self.app.pk]),
                                    {"page-name": "Дахин", "page-path": "/login"})
        self.assertContains(response, "бүртгэлтэй байна")

    def test_page_used_by_scenario_is_not_deleted(self):
        self.client.post(reverse("autotest:page_delete", args=[self.app.pk, self.page.pk]))
        self.assertTrue(Page.objects.filter(pk=self.page.pk).exists())
        self.app.delete()  # апп-ыг устгахад хуудас, сценари хамт устна
        self.assertFalse(Page.objects.exists())

    def test_page_of_other_app_is_rejected(self):
        other_app, *_ = make_setup()
        response = self.client.post(reverse("autotest:scan_create", args=[self.app.pk]),
                                    {"environment": self.env.pk, "page": other_app.pages.get().pk})
        self.assertEqual(response.status_code, 400)
        response = self._post(REGISTER_FIELDS, page=other_app.pages.get().pk)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Scenario.objects.filter(name="Нэвтрэх").exists())

    def test_empty_path_uses_environment_url(self):
        self.env.base_url = "http://web:8000/accounts/login"
        self.env.save()
        home = Page.objects.create(app=self.app, name="Нүүр", path="")
        response = self.client.post(reverse("autotest:scan_create", args=[self.app.pk]),
                                    {"environment": self.env.pk, "page": home.pk})
        self.assertEqual(PageScan.objects.get(pk=response.json()["id"]).url, "http://web:8000/accounts/login")

    def test_scan_is_queued_and_only_owner_can_read_it(self):
        response = self.client.post(reverse("autotest:scan_create", args=[self.app.pk]),
                                    {"environment": self.env.pk, "page": self.page.pk})
        scan = PageScan.objects.get(pk=response.json()["id"])
        self.assertEqual(scan.url, "https://staging.example.com/register")
        scan.status = PageScan.Status.DONE
        scan.result = {"fields": [{"label": "Имэйл", "selector": "#email", "kind": "email"}], "buttons": []}
        scan.save()
        data = self.client.get(reverse("autotest:scan_status", args=[scan.pk]),
                               {"data_file": DataFile.objects.get().pk}).json()
        self.assertEqual((data["fields"][0]["source"], data["fields"][0]["value"]), ("column", "email"))

        self.client.force_login(make_user("qa2", ROLE_QA))
        self.assertEqual(self.client.get(reverse("autotest:scan_status", args=[scan.pk])).status_code, 404)


class RunViewTests(TempMediaMixin, TestCase):
    def setUp(self):
        self.category, _team = make_routed_category()
        self.app, self.env, self.scenario, self.data_file = make_setup(category=self.category)
        self.qa = make_user("qa", ROLE_QA)
        self.client.force_login(self.qa)

    def test_run_is_queued_with_snapshot(self):
        response = self.client.post(reverse("autotest:run_create", args=[self.scenario.pk]),
                                    {"data_file": self.data_file.pk, "environment": self.env.pk})
        run = TestRun.objects.get()
        self.assertRedirects(response, reverse("autotest:run_detail", args=[run.pk]))
        self.assertEqual(run.status, TestRun.Status.QUEUED)
        self.assertEqual(run.target_url, "https://staging.example.com/register")
        self.assertEqual((run.data_file_name, run.environment_name, run.total), ("users", "staging", 1))

    def test_missing_columns_block_the_run(self):
        self.scenario.expected_message_column = "мессеж"
        self.scenario.save()
        self.client.post(reverse("autotest:run_create", args=[self.scenario.pk]),
                         {"data_file": self.data_file.pk, "environment": self.env.pk})
        self.assertFalse(TestRun.objects.exists())

    def _finished_run(self):
        run = TestRun.objects.create(
            scenario=self.scenario, data_file=self.data_file, environment=self.env,
            data_file_name="users", environment_name="staging", target_url="https://staging.example.com/register",
            status=TestRun.Status.DONE, total=2, passed=1, failed=1, started_by=self.qa,
        )
        RunResult.objects.create(run=run, row_number=2, description="Зөв", verdict="pass",
                                 input_data={"email": "a@mail.mn", "password": "***"},
                                 expected_outcome="success", actual_outcome="success")
        failed = RunResult(run=run, row_number=3, description="Давхардсан имэйл", verdict="fail",
                           input_data={"email": "=HYPERLINK(1)", "password": "***"},
                           expected_outcome="error", actual_outcome="success", actual_message="Тавтай морил")
        failed.screenshot.save("shot.png", ContentFile(b"\x89PNG fake"), save=False)
        failed.save()
        return run, failed

    def test_status_json_and_failed_filter(self):
        run, _failed = self._finished_run()
        data = self.client.get(reverse("autotest:run_status", args=[run.pk]), {"show": "failed"}).json()
        self.assertEqual((data["passed"], data["failed"], data["active"]), (1, 1, False))
        self.assertIn("Давхардсан имэйл", data["html"])
        self.assertNotIn("a@mail.mn", data["html"])

    def test_review_rows_are_explained_and_filterable(self):
        run, _failed = self._finished_run()
        RunResult.objects.create(run=run, row_number=4, description="XSS", verdict="recorded",
                                 input_data={"email": "<b>x</b>", "хүлээгдэх": ""},
                                 actual_outcome="error", actual_message="Буруу имэйл")
        run.recorded, run.total = 1, 3
        run.save()
        response = self.client.get(reverse("autotest:run_detail", args=[run.pk]))
        self.assertContains(response, "Гараар шалгах (1)")
        self.assertContains(response, "Дүн юу гэсэн үг вэ?")
        self.assertContains(response, "заагаагүй")
        self.assertNotContains(response, "хүлээгдэх=")  # хүлээгдэх багана оролтод давхар харагдахгүй

        data = self.client.get(reverse("autotest:run_status", args=[run.pk]), {"show": "review"}).json()
        self.assertEqual(data["recorded"], 1)
        self.assertIn("XSS", data["html"])
        self.assertNotIn("Давхардсан имэйл", data["html"])

    def test_export_escapes_formulas(self):
        run, _failed = self._finished_run()
        response = self.client.get(reverse("autotest:run_export", args=[run.pk]))
        sheet = load_workbook(BytesIO(b"".join(response.streaming_content))).worksheets[0]
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        self.assertIn("'=HYPERLINK(1)", values)
        self.assertIn("Унасан", values)

        response = self.client.get(reverse("autotest:run_export", args=[run.pk]), {"format": "csv"})
        _columns, rows = read_rows(BytesIO(b"".join(response.streaming_content)), "r.csv")
        values = [v for _line, row in rows for v in row.values()]
        self.assertIn("'=HYPERLINK(1)", values)
        self.assertIn("Унасан", values)

    def test_bug_ticket_prefilled_then_linked_with_screenshot(self):
        run, failed = self._finished_run()
        url = reverse("autotest:bug_ticket", args=[run.pk]) + f"?ids={failed.pk}"
        response = self.client.get(url)
        self.assertContains(response, "[Автомат тест] Бүртгүүлэх: Давхардсан имэйл")
        self.assertContains(response, "Тавтай морил")

        form = response.context["form"]
        self.assertEqual(form.initial["category"], self.category.pk)  # апп-ын ангиллаар бөглөгдөнө
        response = self.client.post(url, {
            "title": form.initial["title"], "description": form.initial["description"],
            "ticket_type": "bug", "priority": "medium", "category": form.initial["category"],
            "project": make_project().pk,
        })
        ticket = Ticket.objects.get()
        self.assertEqual(ticket.category, self.category)
        self.assertRedirects(response, reverse("tickets:ticket_detail", args=[ticket.pk]))
        failed.refresh_from_db()
        self.assertEqual(failed.ticket, ticket)
        self.assertEqual(ticket.attachments.count(), 1)

    def test_bug_ticket_ignores_passed_rows(self):
        run, _failed = self._finished_run()
        passed = run.results.get(verdict="pass")
        response = self.client.get(reverse("autotest:bug_ticket", args=[run.pk]) + f"?ids={passed.pk}")
        self.assertRedirects(response, reverse("autotest:run_detail", args=[run.pk]))

    def test_cancel_and_rerun(self):
        run, _failed = self._finished_run()
        self.client.post(reverse("autotest:run_rerun", args=[run.pk]))
        new_run = TestRun.objects.exclude(pk=run.pk).get()
        self.assertEqual(new_run.status, TestRun.Status.QUEUED)
        self.client.post(reverse("autotest:run_cancel", args=[new_run.pk]))
        new_run.refresh_from_db()
        self.assertEqual(new_run.status, TestRun.Status.CANCELLED)


class AppCategoryTests(TempMediaMixin, TestCase):
    def setUp(self):
        self.client.force_login(make_user("qa", ROLE_QA))
        self.web, self.mobile = make_category(), make_category()
        self.ios = Subcategory.objects.create(category=self.mobile, name="iOS")

    def test_subcategory_must_belong_to_category(self):
        response = self.client.post(reverse("autotest:app_create"),
                                    {"category": self.web.pk, "subcategory": self.ios.pk, "name": "Shop"})
        self.assertContains(response, "харьяалагдахгүй")
        self.client.post(reverse("autotest:app_create"),
                         {"category": self.mobile.pk, "subcategory": self.ios.pk, "name": "Shop"})
        self.assertEqual(TestApp.objects.get().subcategory, self.ios)

    def test_bug_ticket_gets_app_subcategory(self):
        app, *_ = make_setup(category=self.mobile)
        app.subcategory = self.ios
        app.save()
        run = TestRun.objects.create(scenario=app.scenarios.get(), data_file=DataFile.objects.get(),
                                     environment=app.environments.get(), target_url="https://x", total=1)
        result = RunResult.objects.create(run=run, row_number=2, verdict=RunResult.Verdict.FAIL)
        response = self.client.get(reverse("autotest:bug_ticket", args=[run.pk]) + f"?ids={result.pk}")
        self.assertEqual(response.context["form"].initial["subcategory"], self.ios.pk)

    def test_run_offers_only_files_of_app_category(self):
        app, _env, scenario, own_file = make_setup(category=self.web)
        _other_app, _env2, _s2, other_file = make_setup(category=self.mobile)
        response = self.client.get(reverse("autotest:scenario_detail", args=[scenario.pk]))
        self.assertEqual(list(response.context["run_form"].fields["data_file"].queryset), [own_file])


class PagesRenderTests(TempMediaMixin, TestCase):
    def test_every_page_renders(self):
        app, _env, scenario, data_file = make_setup()
        run = TestRun.objects.create(
            scenario=scenario, data_file=data_file, environment=_env, data_file_name="users",
            environment_name="staging", target_url="https://staging.example.com/register", total=1,
        )
        self.client.force_login(make_user("qa", ROLE_QA))
        for url in [
            reverse("autotest:home"),
            reverse("autotest:app_create"),
            reverse("autotest:app_detail", args=[app.pk]),
            reverse("autotest:datafile_list"),
            reverse("autotest:datafile_create"),
            reverse("autotest:datafile_detail", args=[data_file.pk]),
            reverse("autotest:scenario_create", args=[app.pk]),
            reverse("autotest:scenario_edit", args=[scenario.pk]),
            reverse("autotest:scenario_detail", args=[scenario.pk]),
            reverse("autotest:run_detail", args=[run.pk]),
        ]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)


class OverviewUxTests(TempMediaMixin, TestCase):
    def setUp(self):
        self.app, self.env, self.scenario, self.data_file = make_setup()
        self.client.force_login(make_user("qa", ROLE_QA))

    def test_home_shows_last_run_per_app(self):
        response = self.client.get(reverse("autotest:home"))
        self.assertContains(response, "Одоогоор ажиллуулаагүй")

        TestRun.objects.create(scenario=self.scenario, data_file=self.data_file, environment=self.env,
                               data_file_name="users", environment_name="staging",
                               target_url="https://staging.example.com/register",
                               status=TestRun.Status.DONE, total=4, passed=3, failed=1)
        response = self.client.get(reverse("autotest:home"))
        self.assertContains(response, "/4")
        self.assertEqual(response.context["apps"][0].last_run.passed, 3)

    def test_quick_run_reuses_last_file_and_environment(self):
        response = self.client.get(reverse("autotest:app_detail", args=[self.app.pk]))
        self.assertNotIn("<script>", response.content.decode().split("</title>")[0])
        self.assertNotContains(response, 'name="data_file"')  # өмнө ажиллаагүй — сценари руу оруулна

        last = TestRun.objects.create(scenario=self.scenario, data_file=self.data_file, environment=self.env,
                                      data_file_name="users", environment_name="staging",
                                      target_url="https://staging.example.com/register", status=TestRun.Status.DONE)
        response = self.client.get(reverse("autotest:app_detail", args=[self.app.pk]))
        self.assertContains(response, f'name="data_file" value="{self.data_file.pk}"')
        self.assertContains(response, f'name="environment" value="{self.env.pk}"')

        response = self.client.get(reverse("autotest:scenario_detail", args=[self.scenario.pk]))
        self.assertEqual(response.context["last_run"], last)
        self.assertEqual(response.context["run_form"].initial,
                         {"data_file": self.data_file.pk, "environment": self.env.pk})
        self.assertContains(response, "Тест юу хийх вэ")


class EnvironmentCreateTests(TempMediaMixin, TestCase):
    def test_duplicate_name_keeps_input_and_shows_error(self):
        app, env, *_ = make_setup()
        self.client.force_login(make_user("qa", ROLE_QA))
        url = reverse("autotest:env_create", args=[app.pk])
        response = self.client.post(url, {"env-name": env.name, "env-base_url": "http://web:8000/accounts/login"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f"&#x27;{env.name}&#x27; нэртэй орчин аль хэдийн байна.")
        self.assertContains(response, 'value="http://web:8000/accounts/login"')

        response = self.client.post(url, {"env-name": "prod", "env-base_url": "http://web:8000"})
        self.assertRedirects(response, reverse("autotest:app_detail", args=[app.pk]))
        self.assertEqual(app.environments.count(), 2)

    def test_production_flag_warns_before_running(self):
        app, env, scenario, _file = make_setup()
        self.client.force_login(make_user("qa", ROLE_QA))
        self.client.post(reverse("autotest:env_create", args=[app.pk]),
                         {"env-name": "prod", "env-base_url": "https://shop.mn", "env-is_production": "on"})
        prod = app.environments.get(name="prod")
        self.assertTrue(prod.is_production)
        self.client.post(reverse("autotest:env_toggle_production", args=[app.pk, env.pk]))
        env.refresh_from_db()
        self.assertTrue(env.is_production)

        response = self.client.get(reverse("autotest:scenario_detail", args=[scenario.pk]))
        self.assertEqual(response.context["production_envs"], {env.pk: env.name, prod.pk: "prod"})
        self.assertContains(response, 'data-prod-guard')


class AppLoginTests(TempMediaMixin, TestCase):
    def test_login_settings_are_saved_from_the_test_users_card(self):
        app, *_ = make_setup()
        login_page = Page.objects.create(app=app, name="Нэвтрэх", path="/accounts/login/")
        self.client.force_login(make_user("qa", ROLE_QA))
        url = reverse("autotest:app_login_save", args=[app.pk])
        response = self.client.post(url, {"login-login_page": login_page.pk, "login-api_login_path": "api/login",
                                          "login-api_login_body": "", "login-api_token_prefix": "Bearer"})
        self.assertRedirects(response, reverse("autotest:app_detail", args=[app.pk]))
        app.refresh_from_db()
        self.assertEqual((app.login_page, app.api_login_path), (login_page, "/api/login"))
        # Апп-ын мэдээллийг хадгалахад нэвтрэлт хэвээр.
        self.client.post(reverse("autotest:app_detail", args=[app.pk]),
                         {"name": "Shop2", "category": app.category_id, "description": ""})
        app.refresh_from_db()
        self.assertEqual((app.name, app.login_page), ("Shop2", login_page))


class TestAccountTests(TempMediaMixin, TestCase):
    def setUp(self):
        self.app, self.env, self.scenario, self.data_file = make_setup()
        self.login_page = Page.objects.create(app=self.app, name="Нэвтрэх", path="/accounts/login/")
        self.client.force_login(make_user("qa", ROLE_QA))

    def _save_account(self, password="S3cret!pw", username="qa_test"):
        return self.client.post(reverse("autotest:account_save", args=[self.app.pk]),
                                {"account-label": "QA", "account-username": username, "account-password": password})

    def test_password_is_encrypted_and_never_shown(self):
        self._save_account()
        account = TestAccount.objects.get()
        self.assertNotIn("S3cret!pw", account.password_encrypted)
        self.assertEqual(account.get_password(), "S3cret!pw")
        response = self.client.get(reverse("autotest:app_detail", args=[self.app.pk]))
        self.assertContains(response, "qa_test")
        self.assertNotContains(response, "S3cret!pw")
        self.assertNotContains(response, account.password_encrypted)

    def test_same_label_updates_credentials(self):
        self._save_account()
        self._save_account(password="New!pass1", username="qa2")
        account = TestAccount.objects.get()
        self.assertEqual((account.username, account.get_password()), ("qa2", "New!pass1"))

    def test_scenario_with_account_needs_login_page(self):
        self._save_account()
        account = TestAccount.objects.get()
        data = {"name": "Ticket", "page": self.scenario.page_id, "account": account.pk,
                "success_mode": "auto", "fields_json": json.dumps(REGISTER_FIELDS)}
        response = self.client.post(reverse("autotest:scenario_create", args=[self.app.pk]), data)
        self.assertContains(response, "нэвтрэх хуудсаа сонгоно уу")

        self.app.login_page = self.login_page
        self.app.save()
        self.client.post(reverse("autotest:scenario_create", args=[self.app.pk]), data)
        self.assertEqual(Scenario.objects.get(name="Ticket").account, account)

    def test_scan_and_run_carry_login(self):
        self._save_account()
        account = TestAccount.objects.get()
        self.app.login_page = self.login_page
        self.app.save()
        response = self.client.post(reverse("autotest:scan_create", args=[self.app.pk]),
                                    {"environment": self.env.pk, "page": self.scenario.page_id, "account": account.pk})
        scan = PageScan.objects.get(pk=response.json()["id"])
        self.assertEqual((scan.account, scan.login_url), (account, "https://staging.example.com/accounts/login/"))

        self.scenario.account = account
        self.scenario.save()
        self.client.post(reverse("autotest:run_create", args=[self.scenario.pk]),
                         {"data_file": self.data_file.pk, "environment": self.env.pk})
        run = TestRun.objects.get()
        self.assertEqual((run.account, run.account_label), (account, "QA"))
        self.assertEqual(run.login_url, "https://staging.example.com/accounts/login/")

    def test_account_and_login_page_in_use_are_not_deleted(self):
        self._save_account()
        account = TestAccount.objects.get()
        self.app.login_page = self.login_page
        self.app.save()
        self.scenario.account = account
        self.scenario.save()
        self.client.post(reverse("autotest:account_delete", args=[self.app.pk, account.pk]))
        self.client.post(reverse("autotest:page_delete", args=[self.app.pk, self.login_page.pk]))
        self.assertTrue(TestAccount.objects.exists())
        self.assertTrue(Page.objects.filter(pk=self.login_page.pk).exists())

    def test_access_check_saves_rules_and_queues_run(self):
        self._save_account()
        account = TestAccount.objects.get()
        url = reverse("autotest:access_run_create", args=[self.scenario.pk])
        data = {"environment": self.env.pk, "expect_anon": "denied", f"expect_{account.pk}": "open"}
        self.client.post(url, data)  # нэвтрэх хуудасгүй бол хэрэглэгчээр шалгахгүй
        self.assertFalse(TestRun.objects.exists())

        self.app.login_page = self.login_page
        self.app.save()
        response = self.client.post(url, data)
        run = TestRun.objects.get()
        self.assertRedirects(response, reverse("autotest:run_detail", args=[run.pk]))
        self.assertEqual(run.kind, TestRun.Kind.ACCESS)
        self.assertEqual(run.access_rules, [
            {"account": None, "label": "Нэвтрэхгүй", "expect": "denied"},
            {"account": account.pk, "label": "QA", "expect": "open"},
        ])
        self.assertEqual((run.total, run.login_url), (2, "https://staging.example.com/accounts/login/"))
        self.scenario.refresh_from_db()
        self.assertEqual(self.scenario.access_rules, {"anon": "denied", str(account.pk): "open"})

        detail = self.client.get(reverse("autotest:scenario_detail", args=[self.scenario.pk]))
        self.assertContains(detail, f'name="expect_{account.pk}" value="open" aria-label="QA: Нээлттэй" checked')
        self.client.post(reverse("autotest:run_rerun", args=[run.pk]))
        self.assertEqual(TestRun.objects.filter(kind=TestRun.Kind.ACCESS).count(), 2)

    def test_access_check_is_collapsed_and_defaults_to_scenario_user(self):
        self._save_account()
        account = TestAccount.objects.get()
        self.scenario.account = account
        self.scenario.save()
        detail = self.client.get(reverse("autotest:scenario_detail", args=[self.scenario.pk]))
        self.assertContains(detail, 'id="accessCheck" class="collapse"')
        self.assertContains(detail, f'name="expect_{account.pk}" value="open" aria-label="QA: Нээлттэй" checked')
        self.assertContains(detail, "qa_test")  # "Ажиллуулах" дээр хэн болж шалгахыг харуулна

    def test_access_check_needs_a_choice(self):
        self.client.post(reverse("autotest:access_run_create", args=[self.scenario.pk]),
                         {"environment": self.env.pk, "expect_anon": ""})
        self.assertFalse(TestRun.objects.exists())

    def test_deleted_account_fails_run_with_message(self):
        from apps.autotest.runner import execute_run

        run = TestRun.objects.create(scenario=self.scenario, data_file=self.data_file, environment=self.env,
                                     target_url="https://staging.example.com/register",
                                     login_url="https://staging.example.com/accounts/login/", total=1)
        execute_run(run)
        run.refresh_from_db()
        self.assertEqual(run.status, TestRun.Status.FAILED)
        self.assertIn("устгагдсан", run.error_message)


class RunnerFailureTests(TempMediaMixin, TestCase):
    def test_missing_columns_fail_with_clear_message(self):
        """Файлд сценарийн багана алга бол ойлгомжтой мессежтэй зогсоно (browser нээхгүй)."""
        from apps.autotest.runner import execute_run

        _app, env, scenario, data_file = make_setup(rows=[["Тайлбар", "email", "хүлээгдэх"], ["a", "a@b.mn", ""]])
        run = TestRun.objects.create(scenario=scenario, data_file=data_file, environment=env,
                                     data_file_name="users", environment_name="staging",
                                     target_url=env.url_for(scenario.page.path), total=1)
        execute_run(run)
        run.refresh_from_db()
        self.assertEqual(run.status, TestRun.Status.FAILED)
        self.assertIn("Файлд дараах багана алга: password", run.error_message)


class DataFileMappingTests(TempMediaMixin, TestCase):
    def test_fields_are_remapped_to_new_file_columns(self):
        _app, _env, _scenario, data_file = make_setup(
            rows=[["Тайлбар", "Имэйл", "Нууц үг", "хүлээгдэх"], ["a", "a@b.mn", "x", ""]]
        )
        self.client.force_login(make_user("qa", ROLE_QA))
        fields = [{"label": "Имэйл", "selector": "#email", "kind": "email", "source": "column", "value": "email"},
                  {"label": "Нууц үг", "selector": "#pw", "kind": "password", "source": "column", "value": "password"}]
        data = self.client.post(reverse("autotest:datafile_mapping", args=[data_file.pk]),
                                {"fields_json": json.dumps(fields)}).json()
        self.assertEqual([f["value"] for f in data["fields"]], ["Имэйл", "Нууц үг"])
        self.assertIn("хүлээгдэх", data["columns"])


class DataFilePreviewTests(TempMediaMixin, TestCase):
    def test_preview_returns_rows_as_json(self):
        _app, _env, _scenario, data_file = make_setup()
        self.client.force_login(make_user("qa", ROLE_QA))
        data = self.client.get(reverse("autotest:datafile_preview", args=[data_file.pk])).json()
        self.assertEqual(data["columns"], ["Тайлбар", "email", "password", "хүлээгдэх"])
        self.assertEqual(data["rows"][0], [2, ["Зөв", "a@mail.mn", "Pass1234", "амжилттай"]])
        self.assertEqual(data["total"], 1)
        self.assertTrue(data["roles"]["хүлээгдэх"]["expected"])
        self.assertEqual(data["roles"]["email"]["used_by"], ["Shop · Бүртгүүлэх"])
        self.assertTrue(data["roles"]["password"]["secret"])
        self.assertEqual(data["roles"]["Тайлбар"]["used_by"], [])
        self.assertIn("амжилттай", data["outcome_words"])


class DataFileEditTests(TempMediaMixin, TestCase):
    def setUp(self):
        _app, _env, _scenario, self.data_file = make_setup()
        self.client.force_login(make_user("qa", ROLE_QA))
        self.url = reverse("autotest:datafile_save_rows", args=[self.data_file.pk])

    def _save(self, rows, columns=None):
        payload = {"rows": rows} if columns is None else {"rows": rows, "columns": columns}
        return self.client.post(self.url, json.dumps(payload), content_type="application/json")

    def _read(self):
        self.data_file.refresh_from_db()
        with self.data_file.file.open("rb") as fh:
            return read_rows(fh, self.data_file.file.name)

    def test_unused_columns_can_be_added_renamed_and_removed(self):
        response = self._save(
            [["a@mail.mn", "Pass1", "амжилттай", "x"]], columns=["email", "password", "хүлээгдэх", "Шинэ"],
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["columns"], ["email", "password", "хүлээгдэх", "Шинэ"])
        columns, rows = self._read()
        self.assertEqual(columns, ["email", "password", "хүлээгдэх", "Шинэ"])  # "Тайлбар" устсан
        self.assertEqual(self.data_file.columns, columns)
        self.assertEqual(rows[0][1]["Шинэ"], "x")

    def test_columns_used_by_scenarios_are_kept(self):
        response = self._save([["a", "b", "c", "d"]], columns=["Тайлбар", "e-mail", "password", "хүлээгдэх"])
        self.assertEqual(response.status_code, 400)
        self.assertIn("email", response.json()["error"])
        self.assertEqual(self._save([["a", "b"]], columns=["x", "x"]).status_code, 400)
        self.assertEqual(self._save([["a"]], columns=[" "]).status_code, 400)

    def test_xlsx_keeps_other_sheets_and_styles(self):
        from openpyxl.styles import PatternFill

        with self.data_file.file.open("rb") as fh:
            workbook = load_workbook(fh)
        workbook.worksheets[0]["B2"].fill = PatternFill("solid", fgColor="FFFF00")
        workbook.create_sheet("Тэмдэглэл")["A1"] = "QA notes"
        buffer = BytesIO()
        workbook.save(buffer)
        self.data_file.file.save("users.xlsx", ContentFile(buffer.getvalue()), save=True)

        self._save([["Шинэ", "n@mail.mn", "Pass1", "алдаа"], ["2", "m@mail.mn", "Pass2", "амжилттай"]])
        self.data_file.refresh_from_db()
        with self.data_file.file.open("rb") as fh:
            saved = load_workbook(fh)
        self.assertEqual(saved.sheetnames[1], "Тэмдэглэл")
        self.assertEqual(saved["Тэмдэглэл"]["A1"].value, "QA notes")
        self.assertEqual(saved.worksheets[0]["B2"].fill.fgColor.rgb, "00FFFF00")
        _columns, rows = self._read()
        self.assertEqual([r[1]["email"] for r in rows], ["n@mail.mn", "m@mail.mn"])

    def test_rows_are_saved_and_read_back(self):
        response = self._save([
            ["Зөв", "real@mail.mn", "Real#Pass1", "амжилттай"],
            ["", "", "", ""],  # хоосон мөр хасагдана
            ["Томьёо биш", "=1+1", "x", "алдаа"],
        ])
        self.assertEqual(response.json()["rows"], 2)
        self.data_file.refresh_from_db()
        self.assertEqual(self.data_file.row_count, 2)
        with self.data_file.file.open("rb") as fh:
            columns, rows = read_rows(fh, self.data_file.file.name)
        self.assertEqual(columns, ["Тайлбар", "email", "password", "хүлээгдэх"])
        self.assertEqual(rows[0][1]["email"], "real@mail.mn")
        self.assertEqual(rows[1][1]["email"], "=1+1")  # томьёо болж алга болохгүй

    def test_invalid_rows_rejected(self):
        self.assertEqual(self._save([["only", "two"]]).status_code, 400)
        self.assertEqual(self._save([]).status_code, 400)
        self.assertEqual(self._save("nope").status_code, 400)

    def test_developer_cannot_edit(self):
        self.client.force_login(make_user("dev", ROLE_DEV))
        self._save([["a", "b", "c", "d"]])
        self.data_file.refresh_from_db()
        self.assertEqual(self.data_file.row_count, 1)
